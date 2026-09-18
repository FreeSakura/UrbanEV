"""Convert the complete Markdown manuscript to editable Word via Pandoc OMML."""
from pathlib import Path
import argparse,subprocess,shutil,tempfile,json,zipfile,re,os
from copy import deepcopy
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
PAPER=ROOT/'paper/persistent_events'

def black_style(style,size=None):
    style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
    if size:style.font.size=Pt(size)
    rpr=style.element.get_or_add_rPr();fonts=rpr.find(qn('w:rFonts'))
    if fonts is None:fonts=OxmlElement('w:rFonts');rpr.append(fonts)
    for key in ['ascii','hAnsi']:fonts.set(qn('w:'+key),'Times New Roman')
    fonts.set(qn('w:eastAsia'),'SimSun')

def reference(path,base):
    d=Document(base);sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.8);sec.left_margin=sec.right_margin=Inches(1)
    sec.header_distance=sec.footer_distance=Inches(.35)
    for s in d.styles:
        if s.type in [1,2]:
            try:black_style(s)
            except Exception:pass
        for border in s.element.xpath('.//w:pBdr'):border.getparent().remove(border)
    for name,size in [('Normal',11),('Body Text',11),('Title',18),('Subtitle',12),('Heading 1',13),('Heading 2',11.5),('Caption',9.5)]:
        if name in d.styles:black_style(d.styles[name],size)
    for name in ['Normal','Body Text']:
        f=d.styles[name].paragraph_format;f.line_spacing=1.12;f.space_after=Pt(6);f.widow_control=True
    for name in ['Heading 1','Heading 2']:
        s=d.styles[name];s.font.bold=True;s.paragraph_format.keep_with_next=True
        s.paragraph_format.space_before=Pt(13 if name=='Heading 1' else 9);s.paragraph_format.space_after=Pt(6)
    for name in ['Title','Subtitle']:
        d.styles[name].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('Page ').font.size=Pt(9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
    d.save(path)

def cell_border(table):
    pr=table._tbl.tblPr
    borders=pr.find(qn('w:tblBorders'))
    if borders is None:borders=OxmlElement('w:tblBorders');pr.append(borders)
    for side in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    margins=OxmlElement('w:tblCellMar')
    for side in ['top','bottom','left','right']:
        el=OxmlElement('w:'+side);el.set(qn('w:w'),'85');el.set(qn('w:type'),'dxa');margins.append(el)
    pr.append(margins)

def polish(path):
    d=Document(path)
    for style in d.styles:
        for border in style.element.xpath('.//w:pBdr'):border.getparent().remove(border)
    for sec in d.sections:
        sec.page_width=Inches(8.5);sec.page_height=Inches(11)
        sec.left_margin=sec.right_margin=Inches(1);sec.top_margin=sec.bottom_margin=Inches(.8)
    for p in d.paragraphs:
        p.paragraph_format.widow_control=True
        if p.style.name in ['Title','Subtitle','Author','Date'] or p.text in ['Anonymous manuscript','September 2026']:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True
        if p.style.name.startswith('Heading') or p.style.name in ['Title','Subtitle']:
            for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
            for border in p._p.xpath('.//w:pBdr'):border.getparent().remove(border)
        if p.style.name=='Title':
            for r in p.runs:r.font.size=Pt(18);r.font.name='Times New Roman'
        if p.text.startswith(('Theorem 1','Theorem 2','Corollary 1')):p.paragraph_format.keep_with_next=True
        if p.text.rstrip().endswith('Define'):p.paragraph_format.keep_with_next=True
        if p.text.startswith(('Figure ','Table ')):
            p.style=next(s for s in d.styles if s.style_id=='Caption');p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(8)
            if p.text.startswith('Table '):p.paragraph_format.keep_with_next=True
        if p.text.startswith('[1]') or re.match(r'^\[\d+\] ',p.text):
            p.paragraph_format.left_indent=Inches(.22);p.paragraph_format.first_line_indent=Inches(-.22)
            p.paragraph_format.space_after=Pt(5)
            for r in p.runs:r.font.size=Pt(9.5)
        if p.text.startswith(('Appendix A','Appendix B','中文摘要')):
            p.paragraph_format.page_break_before=True
        if p._p.xpath('.//w:drawing'):
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        if p._p.xpath('.//m:oMathPara'):
            p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(6)
    for shape in d.inline_shapes:
        if shape.width>Inches(6.45):
            ratio=Inches(6.45)/shape.width;shape.width=Inches(6.45);shape.height=int(shape.height*ratio)
    # Rebuild the table containers through python-docx to keep their OOXML
    # schema and paragraph styles portable across Word and LibreOffice.
    for old in list(d.tables):
        new=d.add_table(rows=len(old.rows),cols=len(old.columns))
        for i,row in enumerate(old.rows):
            for j,cell in enumerate(row.cells):
                dest=new.cell(i,j)
                for el in list(dest._tc):
                    if el.tag!=qn('w:tcPr'):dest._tc.remove(el)
                for el in cell._tc:
                    if el.tag!=qn('w:tcPr'):dest._tc.append(deepcopy(el))
        old._tbl.getparent().replace(old._tbl,new._tbl)
    for t in d.tables:
        t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False;cell_border(t)
        n=len(t.columns)
        widths={7:[.95,.65,.70,.55,.4,1.1,1.1],6:[.55,1.55,.60,.75,1.0,1.05],3:[3.65,1.35,1.35]}.get(n,[6.35/n]*n)
        total=sum(widths);widths=[w*6.35/total for w in widths]
        for j,col in enumerate(t.columns):col.width=Inches(widths[j])
        for i,row in enumerate(t.rows):
            if i==0:
                repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
            no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
            for j,cell in enumerate(row.cells):
                cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF1' if i==0 else 'FFFFFF');tcpr.append(shade)
                for p in cell.paragraphs:
                    p.style=d.styles['Normal'];p.paragraph_format.keep_with_next=False
                    p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.0
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or (n==6 and j==1) else WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor(0,0,0);r.bold=i==0
    d.core_properties.title='Persistent Event Evaluation under Partial Observations'
    d.core_properties.author='Anonymous';d.core_properties.last_modified_by='';d.core_properties.comments=''
    mathpr=OxmlElement('m:mathPr');mathfont=OxmlElement('m:mathFont');mathfont.set(qn('m:val'),'Cambria Math');mathpr.append(mathfont);d.settings.element.append(mathpr)
    # Retain standard OMML script styles; supplementary Unicode math letters
    # can be dropped by some LibreOffice importers.
    d.save(path)

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--pandoc',default=shutil.which('pandoc'));ap.add_argument('--output',type=Path,required=True)
    ap.add_argument('--resource-root',type=Path,default=PAPER,help='Parent of figures/ when building with temporary figures')
    a=ap.parse_args()
    if not a.pandoc:ap.error('Install Pandoc or supply --pandoc')
    a.output.parent.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='persistent-word-') as tmp:
        base=Path(tmp)/'base.docx'
        base.write_bytes(subprocess.run([str(a.pandoc),'--print-default-data-file','reference.docx'],capture_output=True,check=True).stdout)
        ref=Path(tmp)/'reference.docx';reference(ref,base)
        source=Path(tmp)/'manuscript.md'
        source.write_text(re.sub(r'\\tag\{(\d+)\}',lambda m:r'\quad\text{('+m.group(1)+')}',(PAPER/'manuscript.md').read_text(encoding='utf-8')),encoding='utf-8')
        resource_path=os.pathsep.join([str(a.resource_root.resolve()),str(PAPER)])
        cmd=[str(a.pandoc),str(source),'--from=markdown-implicit_figures','--standalone','--reference-doc='+str(ref),'--resource-path='+resource_path,'--output='+str(a.output)]
        r=subprocess.run(cmd,text=True,capture_output=True,check=True)
        if r.stderr:print(r.stderr)
        if 'Could not convert' in r.stderr:raise RuntimeError('Equation conversion failed; raw LaTeX is not an acceptable Word deliverable')
        polish(a.output)
    with zipfile.ZipFile(a.output) as z:
        xml=z.read('word/document.xml').decode('utf-8')
        count=len(re.findall(r'<m:oMath(?:\s|>)',xml))
        assert count>100,'Expected editable native equations in this mathematical manuscript'
        assert 'documentProtection' not in z.read('word/settings.xml').decode('utf-8')
    print(json.dumps({'output':str(a.output),'native_math_objects':count,'text_and_tables_editable':True}))

if __name__=='__main__':main()
